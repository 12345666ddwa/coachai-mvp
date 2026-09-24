#!/usr/bin/env python3
"""Generate the all-English direction confirmation document for the team."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

HEADER_BG = "1F4E79"
ALT_ROW = "F2F6FA"
DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x6B, 0x76, 0x84)
RED = RGBColor(0xC0, 0x39, 0x2B)

def set_cell_shading(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def set_cell_text(cell, text, bold=False, color=None, size=Pt(9.5)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run.font.size = size
    run.bold = bold
    if color:
        run.font.color.rgb = color

def create_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, HEADER_BG)
        set_cell_text(cell, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_text(cell, str(val))
            if r_idx % 2 == 0:
                set_cell_shading(cell, ALT_ROW)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Calibri"; r.font.size = Pt(14); r.bold = True
    r.font.color.rgb = DARK_BLUE

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Calibri"; r.font.size = Pt(11.5); r.bold = True

def body(doc, text, size=Pt(10.5)):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Calibri"; r.font.size = size

def bullet(doc, text, size=Pt(10.5), bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5)
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        r0.font.name = "Calibri"; r0.font.size = size; r0.bold = True
    r = p.add_run(text)
    r.font.name = "Calibri"; r.font.size = size

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

# ---------- Title ----------
p = doc.add_paragraph()
r = p.add_run("CoachAI - Next Phase: Direction & Scope Proposal")
r.font.name = "Calibri"; r.font.size = Pt(20); r.bold = True; r.font.color.rgb = DARK_BLUE
p = doc.add_paragraph()
r = p.add_run("For team review and confirmation | prepared by Xing (Happy)")
r.font.name = "Calibri"; r.font.size = Pt(11); r.font.color.rgb = GRAY
body(doc, "")

# ---------- 1. Why ----------
h1(doc, "1. Why this document")
body(doc, "We are about to start the next phase of development. Before writing any code, we want to confirm "
          "with the team that the planned features, the intended outcomes, and the design direction match "
          "everyone's expectations.")
body(doc, "Please read this document, add comments or edits directly, and reply with either a confirmation or "
          "change requests. We will not start building until the direction is agreed. Once agreed, we will "
          "work through the plan in section 6 and report progress at each milestone.")

# ---------- 2. Where we are ----------
h1(doc, "2. Where we are today (recap)")
bullet(doc, "dual-agent marking engine (Marker + Verifier), confidence scoring, teacher-in-the-loop "
            "flagging - all running on 18 real 2025 HSC questions", bold_prefix="Working prototype: ")
bullet(doc, "32-answer regression against the official NESA rubric: 78.1% exact match, 100% within "
            "+/-1 band, zero wild misses", bold_prefix="Verified accuracy: ")
bullet(doc, "live AI demo + static concept site + full progress report (links shared in chat)", bold_prefix="Live assets: ")
bullet(doc, "full question display fixed, cleaner interface (no emoji, squarer components) - delivered in "
            "the last iteration", bold_prefix="Team feedback already actioned: ")

# ---------- 3. Proposed features ----------
h1(doc, "3. What we propose to build next")
body(doc, "Four workstreams. Each is described as: what it is, why it matters, how it works, and what you "
          "will see in the demo.")

h2(doc, "3.1  Privacy: local anonymisation layer")
bullet(doc, "a local anonymiser that replaces student names, IDs and school identifiers with placeholders "
            "(e.g. [STUDENT_A]) BEFORE any content is sent to the AI. Real values are only re-attached "
            "for display.", bold_prefix="What: ")
bullet(doc, "addresses the team requirement that student data must be anonymous - \"ok for demo but needs "
            "to be fully private\". It also directly answers the privacy question any judge will ask.",
       bold_prefix="Why: ")
bullet(doc, "runs entirely on the local machine; nothing identifiable leaves the device. For the demo we "
            "can show a side-by-side: what the teacher sees vs. what the AI actually receives.",
       bold_prefix="How: ")
bullet(doc, "in the demo, type an answer containing a name -> the AI-side transcript shows "
            "[STUDENT_A]; the result card still displays the real name.", bold_prefix="What you will see: ")

h2(doc, "3.2  AI Lesson Planner (new feature)")
bullet(doc, "a lesson-plan generator driven by a checkbox interface. The teacher imports reference "
            "material, ticks the topics they want covered (and what has already been taught), and the AI "
            "produces a structured lesson plan.", bold_prefix="What: ")
bullet(doc, "extends the product from \"marking assistant\" to \"teaching assistant\", covering the "
            "planning side of a teacher's workload.", bold_prefix="Why: ")
bullet(doc, "the same NESA Teacher Support Resources that power our marking RAG also feed the lesson "
            "planner - no new data pipeline needed. Output is structured: learning objectives, sequence "
            "of activities, assessment points, and timing.", bold_prefix="How: ")
bullet(doc, "select topics -> import a document's text -> a generated lesson plan appears, ready to "
            "copy into the teacher's own template.", bold_prefix="What you will see: ")
bullet(doc, "the reference document will be emailed to us. Meanwhile we will build the interface and "
            "engine using the official NESA TSR materials we already have.", bold_prefix="Dependency: ")

h2(doc, "3.3  Marking upgrades (from the team's task notes)")
bullet(doc, "the AI will phrase results as a suggested mark / draft evaluation the teacher can accept or "
            "adjust - not a confident final judgement.", bold_prefix="Tone change: ")
bullet(doc, "each judgement will explicitly reference the exact marking-criteria line(s) it was based "
            "on, so the teacher can see the reasoning at a glance.", bold_prefix="Criteria transparency: ")
bullet(doc, "layout and whitespace polish, aligned with the team's Web Dev notes (see section 5 for the "
            "one open question).", bold_prefix="Visual cleanup: ")

h2(doc, "3.4  Interface redesign (anti-AI-coded)")
bullet(doc, "the team noted the interface should not look AI-coded. We have compiled the full \"things to "
            "avoid\" list (no emoji, no curved boxes, no gradients, no glassmorphism, no decorative "
            "animation, no template fonts) and will design directly against it.", bold_prefix="Direction: ")
bullet(doc, "we propose an exam-paper design language: warm paper background, ink text, and a single "
            "red-pen accent - the visual metaphor of a teacher's marked paper. Distinctive, "
            "education-native, and unlike any generic template.", bold_prefix="Proposed visual identity: ")

h2(doc, "3.5  Report generation (optional - depends on scope decision)")
bullet(doc, "generate NESA-aligned report comments from a spreadsheet of student results, using the "
            "report-comment guidance already added to the shared document.", bold_prefix="What: ")
bullet(doc, "only included if the team picks scope option C below; it would follow the core workstreams.",
       bold_prefix="Scope: ")

# ---------- 4. Scope options ----------
h1(doc, "4. Scope options - please choose one")
create_table(doc, ["Option", "Contents", "Risk", "Recommendation"], [
    ["A - Focused addition",
     "Keep everything that is built and verified. Add the four workstreams in section 3 (3.1-3.4).",
     "Low", "Recommended"],
    ["C - A + report prototype",
     "All of option A, plus 3.5 report generation as a working prototype.",
     "Medium", "Good if time allows"],
    ["B - Full rebuild",
     "Rebuild the product end-to-end around the new features, dropping current assets.",
     "High - throws away verified work", "Not recommended"],
], col_widths=[3.4, 7.6, 3.2, 3.0])

# ---------- 5. Design checklist ----------
h1(doc, "5. Design principles (the anti-AI-coded checklist)")
body(doc, "We will check every screen against this list before showing anything. It merges the team's "
          "notes with the \"things to avoid\" list:")
create_table(doc, ["Avoid", "We will instead"], [
    ["Purple-blue gradients, gradient hero text", "Solid colours only: paper, ink, one red accent"],
    ["Emoji in headings, icons everywhere", "No emoji; text and hairline rules instead of icons"],
    ["Glassmorphism / coloured-border cards / 3-icon rows", "Exam-paper layout: lines, spacing, typographic hierarchy"],
    ["Low-contrast dark mode", "Warm light theme (readable, calibrated contrast)"],
    ["Badge-above-headline, cursor-follow beams, scroll fades", "Motion only for real state changes (e.g. marking progress)"],
    ["Template fonts (Inter everywhere, trendy serif pairs)", "System font stack with deliberate size/weight hierarchy"],
    ["Em dashes, generic buzzword copy", "Plain, specific language; no em dashes"],
], col_widths=[7.6, 8.6])

# ---------- 6. Plan ----------
h1(doc, "6. Execution plan (after your confirmation)")
create_table(doc, ["Phase", "Content", "Output"], [
    ["1. Design", "Two HTML design mockups: marking page + lesson planner page, checked against the section 5 list", "Link to mockups for your review"],
    ["2. Privacy", "Anonymisation layer + visible redaction demo", "Working demo"],
    ["3. Lesson planner", "Checkbox UI + document import + plan generation", "Working demo"],
    ["4. Marking upgrades", "Suggested-mark phrasing + criteria references + layout polish", "Updated demo"],
    ["5. Integration & QA", "Full walkthrough, team test round, final deployment", "Final demo + updated report"],
], col_widths=[3.6, 9.0, 3.6])

# ---------- 7. Open questions ----------
h1(doc, "7. Open questions for the team")
create_table(doc, ["#", "Question", "Our suggestion"], [
    ["1", "Which scope - A or C?", "A as the base, C if the schedule allows"],
    ["2", "Reference document for the lesson planner - when will it be emailed?", "Build with NESA TSR materials in the meantime"],
    ["3", "Web Dev notes: \"remove whitespace\" vs \"add white space\" - which do you want?", "Remove dead space, keep breathing room around content (both, done deliberately)"],
    ["4", "Can anyone share a real (anonymised) sample of past student work for testing?", "Otherwise we continue with our 32-answer test set"],
    ["5", "Who will supply the lesson-topic checklist (the tickable list)?", "We can draft one from the syllabus and you review it"],
], col_widths=[0.9, 9.3, 6.0])

# ---------- 8. What we need ----------
h1(doc, "8. What we need from you")
bullet(doc, "Read this document; comment or edit directly, or reply in the chat.", bold_prefix="1. ")
bullet(doc, "Confirm the scope (A or C) and answer the five questions in section 7.", bold_prefix="2. ")
bullet(doc, "Once confirmed, we start Phase 1 and share the design mockups before any code is written.",
       bold_prefix="3. ")
body(doc, "")
body(doc, "We will not start building until we hear back - please take a few minutes to review. "
          "Thank you!")

OUT = "/home/gaogao/workspace/ai-coach/CoachAI_NextPhase_Proposal_EN.docx"
doc.save(OUT)
print("saved:", OUT)
