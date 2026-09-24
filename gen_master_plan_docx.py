#!/usr/bin/env python3
"""Convert MASTER_WORK_PLAN_v1.md to a formatted docx (CN)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml
from docx.enum.table import WD_TABLE_ALIGNMENT

HEADER_BG = "1F4E79"
ALT_ROW = "F2F6FA"
DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x6B, 0x76, 0x84)

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)

def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def styled_run(p, text, size, bold=False, color=None, mono=False, font="微软雅黑"):
    r = p.add_run(text)
    r.font.name = "Consolas" if mono else font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if not bold else "黑体")
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = color
    return r

lines = open('/home/gaogao/workspace/ai-coach/MASTER_WORK_PLAN_v1.md', encoding='utf-8').read().split('\n')

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    # skip empty
    if not line.strip():
        i += 1
        continue
    # code fence
    if line.strip().startswith('```'):
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i])
            i += 1
        i += 1
        for cl in code_lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            styled_run(p, cl if cl else ' ', 9, mono=True, font="Consolas")
        continue
    # table
    if line.strip().startswith('|'):
        tbl_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            tbl_lines.append(lines[i].strip())
            i += 1
        rows = []
        for tl in tbl_lines:
            if set(tl.replace('|', '').replace('-', '').replace(' ', '')) == set():
                continue
            cells = [c.strip() for c in tl.strip('|').split('|')]
            rows.append(cells)
        if rows:
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = table.rows[ri].cells[ci]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    val = row[ci] if ci < len(row) else ''
                    if ri == 0:
                        shade(cell, HEADER_BG)
                        styled_run(p, val, 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                    else:
                        if ri % 2 == 0:
                            shade(cell, ALT_ROW)
                        styled_run(p, val, 9)
            doc.add_paragraph()
        continue
    # headings
    if line.startswith('# '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(8)
        styled_run(p, line[2:], 17, bold=True, color=DARK_BLUE)
    elif line.startswith('## '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
        styled_run(p, line[3:], 14, bold=True, color=DARK_BLUE)
    elif line.startswith('### '):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
        styled_run(p, line[4:], 12, bold=True)
    elif line.startswith('**') and line.endswith('**') and len(line) < 120:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        styled_run(p, line.strip('*'), 11, bold=True)
    elif line.startswith('- [ ] ') or line.startswith('- [x] '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(2)
        styled_run(p, ('☐ ' if '[ ]' in line[:6] else '☑ ') + line[6:], 10)
    elif line.startswith('- '):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6); p.paragraph_format.space_after = Pt(2)
        styled_run(p, '• ' + line[2:], 10)
    elif line.startswith('---'):
        pass
    else:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(4)
        # bold segments **...**
        import re
        parts = re.split(r'(\*\*[^*]+\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                styled_run(p, part[2:-2], 10.5, bold=True)
            else:
                styled_run(p, part, 10.5)
    i += 1

OUT = "/home/gaogao/workspace/ai-coach/工作总纲_v1.0.docx"
doc.save(OUT)
print("saved:", OUT)
