#!/usr/bin/env python3
"""Generate: 1) full transcript docx (EN), 2) meeting minutes EN, 3) meeting minutes CN."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

HEADER_BG = "1F4E79"
ALT_ROW = "F2F6FA"
DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x6B, 0x76, 0x84)

def set_cell_shading(cell, color):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def set_cell_text(cell, text, bold=False, color=None, size=Pt(10)):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run.font.size = size
    run.bold = bold
    if color:
        run.font.color.rgb = color

def create_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, HEADER_BG)
        set_cell_text(cell, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_text(cell, str(val))
            if r_idx % 2 == 0:
                set_cell_shading(cell, ALT_ROW)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table

def h1(doc, text, font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(15); r.bold = True
    r.font.color.rgb = DARK_BLUE

def h2(doc, text, font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(12); r.bold = True

def body(doc, text, size=Pt(10.5), font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = size

def bullet(doc, text, size=Pt(10.5), font="Calibri"):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = size

def cover(doc, title, sub, meta):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.name = "Calibri"; r.font.size = Pt(28); r.bold = True; r.font.color.rgb = DARK_BLUE
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub); r.font.name = "Calibri"; r.font.size = Pt(13)
    doc.add_paragraph()
    for m in meta:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(m); r.font.name = "Calibri"; r.font.size = Pt(11)
    doc.add_page_break()

OUT_DIR = "/home/gaogao/workspace/ai-coach/meeting/"

# ============================================================ 1. FULL TRANSCRIPT
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)
cover(doc, "Meeting Recording - Full Transcript",
      "Datathon pre-competition workshop (technical training + competition briefing)",
      ["Source: 2026-09-04 18:13 recording (34 min, auto-transcribed)", "Language: English (auto-detected)", "Transcript contains ASR inaccuracies - speaker names and terms marked where unclear"])

import re
with open("/home/gaogao/workspace/ai-coach/meeting_transcript.txt", encoding="utf-8") as f:
    lines = f.read().split("\n")
for ln in lines:
    if ln.startswith("#"):
        continue
    m = re.match(r"\[(\d+:\d+-\d+:\d+)\] (.*)", ln)
    if m:
        ts = m.group(1)
        text = m.group(2)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(f"[{ts}] ")
        r1.font.name = "Consolas"; r1.font.size = Pt(8); r1.font.color.rgb = GRAY
        r2 = p.add_run(text)
        r2.font.name = "Calibri"; r2.font.size = Pt(10)
doc.save(OUT_DIR + "Transcript_Full_2026-09-04.docx")
print("transcript docx saved")

# ============================================================ 2. MINUTES EN
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)
cover(doc, "Workshop & Briefing Notes",
      "Datathon kickoff workshop - 4 September 2026",
      ["Session: technical training workshop + competition Q&A",
       "Speakers: Lucas (academic instructor), Lillian (UNSW), Edge Directors",
       "Prepared for the team | source: full recording transcript"])

h1(doc, "1. Session Overview")
body(doc, "A pre-competition workshop designed to give participants a starting point before the "
          "datathon begins. It covers practical data-science tools in four stages, followed by a "
          "Q&A with the key competition logistics. Note: the audio was auto-transcribed; a few "
          "names/terms are unclear (organiser name, competition codename, and platform names) - "
          "marked where relevant.")

h1(doc, "2. Technical Workshop - Four Stages")
h2(doc, "Stage 1: Understand your data")
bullet(doc, "Import the dataset; identify variable types; inspect columns and dtypes.")
bullet(doc, "pandas essentials: df.head(), df.info() (overview + missing values), df.describe() "
            "(descriptive stats), df.shape, df.dtypes.")
bullet(doc, "Selecting data: df['col'], .loc for rows/columns, boolean masks for conditional "
            "filtering (e.g. observations where x > 5).")
h2(doc, "Stage 2: Clean and prepare data")
bullet(doc, "Missing values: drop rows (.dropna), or impute with mean / median. Median is safer "
            "than mean for skewed data. IMPORTANT: do not assume missingness is meaningless - the "
            "fact a value is missing can itself carry information (example given: income).")
bullet(doc, "Categorical variables: ML models need numeric input - use one-hot / binary encoding, "
            "or ordinal encoding when categories have a natural order (e.g. education level).")
bullet(doc, "Outliers and scale: consider transformations (log for skewed, square root for "
            "moderately skewed) and scaling (e.g. 0-1 vs 0-1000 ranges).")
bullet(doc, "Useful tools: .groupby(), .value_counts(), .merge(), crosstab() for categorical "
            "combinations, duplicated()/drop_duplicates().")
h2(doc, "Stage 3: Explore and model")
bullet(doc, "Visualise patterns and relationships (e.g. correlation matrix) before modelling.")
bullet(doc, "Correlation does not imply causation - treat it as a hint, not proof.")
bullet(doc, "Feature engineering, then train a simple baseline model first; compare everything "
            "against it.")
h2(doc, "Stage 4: Evaluate and submit")
bullet(doc, "Workflow: training set -> develop model -> submit predictions -> evaluated against the "
            "competition metric (on the platform, referred to as 'cable' in the audio - likely Kaggle).")
bullet(doc, "Make sure the submission file is formatted correctly.")

h1(doc, "3. General Advice from Speakers")
bullet(doc, "Understand your data before modelling - explore distributions and relationships first.")
bullet(doc, "Start simple: a solid baseline beats a rushed complex model; use it as the comparison point.")
bullet(doc, "Experiment: try different features, preprocessing and models, and keep track of changes.")
bullet(doc, "Research: the workshop is only a foundation - plenty of resources online.")

h1(doc, "4. Competition Logistics (from Q&A)")
h2(doc, "Timeline (relative to the recording on Friday 4 September)")
create_table(doc, ["When", "What"], [
    ["Tomorrow (opening ceremony, 10:00-12:00)", "Datasets + problem statement released after the ceremony; 3 streams/tracks announced; team registration"],
    ["Sunday 23:59", "Submission deadline: predictions + project submitted online (Kaggle-style)"],
    ["Monday (closing ceremony, 17:00-19:00)", "Top teams invited to judging round - live presentation to judges/sponsors; winners announced"],
], col_widths=[6.0, 10.0])
h2(doc, "Teams")
bullet(doc, "Team size: 1 to 4 people (individual participation allowed).")
bullet(doc, "All team members must be from the SAME state (live presentation judging makes "
            "cross-state teams impractical).")
bullet(doc, "Form teams yourself via the Discord team-finding channel, or register solo and be "
            "randomly assigned - assignments will be posted on Discord.")
bullet(doc, "Use an email you check regularly when registering (used to contact your group).")
h2(doc, "Communication")
bullet(doc, "Discord = PRIMARY channel: Q&A, announcements, team forming, workshop follow-ups.")
bullet(doc, "Instagram (organisers) = secondary updates; the session will also be streamed/recorded "
            "(Zoom) for later viewing.")
bullet(doc, "Questions not answered today will be addressed at the opening ceremony and posted on Discord.")

h1(doc, "5. Action Items for Us")
create_table(doc, ["#", "Action", "Owner / Note"], [
    ["1", "Join the competition Discord (most important)", "Everyone - check the invite link"],
    ["2", "Follow the organisers on Instagram", "Everyone"],
    ["3", "Register / attend the opening ceremony tomorrow 10:00-12:00", "Team"],
    ["4", "Decide team composition (1-4 people, same state) before/at registration", "Team"],
    ["5", "Watch the workshop recording if anything was missed", "Everyone"],
    ["6", "Prepare for the Sunday 23:59 submission once the dataset drops", "Team"],
], col_widths=[1.0, 9.5, 5.5])

h1(doc, "6. Open Questions / To Confirm")
bullet(doc, "Exact platform name and submission mechanics (released tomorrow).")
bullet(doc, "What the 3 streams are and how to pick one.")
bullet(doc, "Judging criteria for the presentation round (Monday).")
doc.save(OUT_DIR + "Meeting_Minutes_EN_2026-09-04.docx")
print("minutes EN saved")

# ============================================================ 3. MINUTES CN
doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2); s.bottom_margin = Cm(2); s.left_margin = Cm(2); s.right_margin = Cm(2)
cover(doc, "赛前工作坊会议纪要",
      "Datathon 开幕前训练营（技术培训 + 比赛说明）",
      ["日期：2026年9月4日（录音 34 分钟）",
       "主讲：Lucas（学术讲师）、Lillian（UNSW）、Edge Directors",
       "整理：高兴（基于完整录音转写）"])

h1(doc, "一、会议性质")
body(doc, "比赛开始前的培训工作坊，帮参赛者建立数据科学基础技能，随后是比赛规则的 Q&A。"
          "录音为自动转写，部分人名/术语听写不准（组织方名称、比赛代号、平台名等），已标注待确认。", font="宋体")

h1(doc, "二、技术工作坊：四阶段框架")
h2(doc, "阶段 1：理解数据", font="黑体")
bullet(doc, "导入数据集、识别变量类型、查看列和数据类型。", font="宋体")
bullet(doc, "pandas 基础：df.head()、df.info()（概览+缺失值）、df.describe()（描述统计）、df.shape、df.dtypes。", font="宋体")
bullet(doc, "数据选取：df['列名']、.loc 选行/列、布尔掩码做条件筛选（如 x>5 的观测）。", font="宋体")
h2(doc, "阶段 2：清洗与准备数据", font="黑体")
bullet(doc, "缺失值：删除（dropna）或填充均值/中位数——偏态数据用中位数更稳。⚠️ 不要默认缺失无意义：缺失本身可能携带信息（举例：收入字段缺失）。", font="宋体")
bullet(doc, "类别变量：模型只吃数值——用 one-hot/哑变量编码；有自然顺序的（如学历）用序数编码。", font="宋体")
bullet(doc, "异常值与量纲：偏态用对数变换、中度偏态用平方根；注意数值范围差异（0-1 vs 0-1000）需归一化。", font="宋体")
bullet(doc, "常用工具：groupby、value_counts、merge、crosstab（类别组合）、duplicated/drop_duplicates。", font="宋体")
h2(doc, "阶段 3：探索与建模", font="黑体")
bullet(doc, "先可视化找规律（如相关性矩阵）。相关≠因果，只能当线索。", font="宋体")
bullet(doc, "特征工程后，先跑一个简单基线模型，后续一切对比它。", font="宋体")
h2(doc, "阶段 4：评估与提交", font="宋体")
bullet(doc, "流程：训练集 → 建模 → 提交预测 → 平台按官方 metric 自动评分（录音里的平台名听作 cable，应为 Kaggle 一类）。", font="宋体")
bullet(doc, "提交文件格式必须正确。", font="宋体")

h1(doc, "三、主讲人通用建议", font="黑体")
bullet(doc, "建模前先吃透数据——先探索分布和关系。", font="宋体")
bullet(doc, "从简开始：扎实的基线胜过仓促的复杂模型。", font="宋体")
bullet(doc, "多实验：换特征/预处理/模型，并记录每次改动。", font="宋体")
bullet(doc, "多查资料：工作坊只是起点，网上资源很多。", font="宋体")

h1(doc, "四、比赛关键信息（Q&A 部分）", font="黑体")
h2(doc, "时间线（录音为 9月4日 周五）", font="黑体")
create_table(doc, ["时间", "事项"], [
    ["明天（周六）开幕式 10:00-12:00", "仪式后发布数据集+问题陈述；公布 3 个赛道（streams）；组队登记"],
    ["周日 23:59", "提交截止：预测结果 + 项目（Kaggle 式线上提交）"],
    ["周一（闭幕式 17:00-19:00）", "晋级队伍现场向评委/赞助商演示；颁奖"],
], col_widths=[6.0, 10.0])
h2(doc, "组队规则", font="黑体")
bullet(doc, "人数：1-4 人（允许个人参赛）。", font="宋体")
bullet(doc, "所有队员必须在同一州（因为晋级后要现场演示，跨州不现实）。", font="宋体")
bullet(doc, "可在 Discord 找队友自行组队；也可单独报名等随机分配，分配结果发 Discord。", font="宋体")
bullet(doc, "报名邮箱填常用邮箱（用于联系组员）。", font="宋体")
h2(doc, "沟通渠道", font="宋体")
bullet(doc, "Discord = 第一渠道：Q&A、公告、组队、工作坊后续。", font="宋体")
bullet(doc, "Instagram（组织方）= 次要；工作坊有 Zoom 录播可回看。", font="宋体")
bullet(doc, "今天没答完的问题：开幕式再答 + Discord 上发。", font="宋体")

h1(doc, "五、我们的行动清单", font="黑体")
create_table(doc, ["#", "行动", "负责人/备注"], [
    ["1", "加入比赛 Discord（最重要）", "全员——找邀请链接"],
    ["2", "关注组织方 Instagram", "全员"],
    ["3", "报名并参加明天开幕式 10:00-12:00", "团队"],
    ["4", "定组队（1-4 人、同州）", "团队"],
    ["5", "有遗漏的看工作坊录播", "全员"],
    ["6", "数据集发布后备战周日 23:59 提交", "团队"],
], col_widths=[1.0, 9.5, 5.5])

h1(doc, "六、待确认问题", font="黑体")
bullet(doc, "提交平台与具体操作（明天发布）。", font="宋体")
bullet(doc, "3 个赛道是什么、怎么选。", font="宋体")
bullet(doc, "周一演示轮的评分标准。", font="宋体")
doc.save(OUT_DIR + "会议纪要_中文_2026-09-04.docx")
print("minutes CN saved")
